"""文件连接器：csv / json / xlsx → list[dict]（传输机制，域无关）。

上迁源：APS ``platform/connectors/file.py``。

依赖纪律：csv/json 走 stdlib；xlsx 优先 pandas，缺失回退 openpyxl
（两者皆缺 ⇒ 明确 ImportError，不静默跳过格式）。
输出契约：``rows`` 为 ``list[dict]``（首行/表头为键），**不做换算/归一**
（归采纳期，``no_conversion_in_connector``）。
"""
from __future__ import annotations

import csv
import json
import os
from typing import Any, Dict, List, Mapping

from .base import ConnectorResult, SourceConnector

__all__ = ["FileConnector"]


class FileConnector(SourceConnector):
    """本地文件源连接器（csv / json / xlsx）。"""

    name = "file"

    def fetch(self, cfg: Mapping[str, Any]) -> ConnectorResult:
        path = cfg.get("path")
        if not path:
            raise ValueError("FileConnector 需要 cfg['path']")
        if not os.path.exists(path):
            raise FileNotFoundError(f"文件不存在：{path}")

        ext = os.path.splitext(path)[1].lower().lstrip(".")
        if ext == "csv":
            rows = self._read_csv(path)
        elif ext == "json":
            rows = self._read_json(path)
        elif ext in ("xlsx", "xls"):
            rows = self._read_xlsx(path)
        else:
            raise ValueError(f"不支持的文件类型：{ext!r}（仅 csv/json/xlsx）")

        return ConnectorResult(
            rows=rows,
            external_refs=[path],
            meta={"ext": ext, "n_rows": len(rows)},
        )

    # -- 各格式读取（纯函数，便于单测）-------------------------------------

    @staticmethod
    def _read_csv(path: str) -> List[Dict[str, Any]]:
        # utf-8-sig 兼容带 BOM 的 Excel 导出
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [dict(r) for r in csv.DictReader(f)]

    @staticmethod
    def _read_json(path: str) -> List[Dict[str, Any]]:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return [dict(r) for r in data]
        if isinstance(data, dict) and "rows" in data:
            return [dict(r) for r in data["rows"]]
        raise ValueError("JSON 须为 list[dict] 或 {'rows': [...]}")

    @staticmethod
    def _read_xlsx(path: str) -> List[Dict[str, Any]]:
        try:
            import pandas as pd
        except ImportError:
            return FileConnector._read_xlsx_openpyxl(path)
        df = pd.read_excel(path)
        # NaN → None，避免采纳期收到 float('nan') 污染外部字段
        return df.where(df.notna(), None).to_dict(orient="records")

    @staticmethod
    def _read_xlsx_openpyxl(path: str) -> List[Dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - 环境无两者时
            raise ImportError(
                "读取 xlsx 需要 pandas 或 openpyxl（皆缺失）"
            ) from exc

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(h) for h in next(it)]
        except StopIteration:
            return []
        out: List[Dict[str, Any]] = []
        for r in it:
            if all(c is None for c in r):
                continue
            out.append({k: v for k, v in zip(header, r)})
        return out
